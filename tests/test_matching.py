"""Self-checks for resume parsing, JD parsing and matching (no Streamlit or internet needed).

Run from the project root:   python tests/test_matching.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from modules.document_parser import extract_text
from modules.jd_parser import parse_job_description
from modules.resume_parser import parse_resume
from modules.similarity_engine import compare_resume_to_jd
from modules.skill_extractor import extract_skills, match_skill
from modules.skill_gap import build_skill_gap

ROOT = Path(__file__).resolve().parent.parent

# ---------------------------------------------------------------- skills
assert extract_skills("Java and JavaScript") == ["Java", "JavaScript"]          # Java is not found inside JavaScript
assert "C++" in extract_skills("C++, C#") and "C#" in extract_skills("C++, C#")
assert match_skill("SQL", {"MySQL"})[0] == "matched"                              # MySQL proves SQL
assert match_skill("Machine Learning", {"TensorFlow"})[0] == "matched"
assert match_skill("Azure", {"AWS"})[0] == "partial"                              # other cloud: partial only
assert match_skill("Docker", {"Python"})[0] == "missing"

# ---------------------------------------------------------------- job description
jd = parse_job_description((ROOT / "sample_data" / "sample_job_description.txt").read_text(encoding="utf-8"))
assert jd["required_skills"] == ["Data Analysis", "Git", "Machine Learning", "Pandas", "Python", "SQL"], jd["required_skills"]
assert jd["preferred_skills"] == ["AWS", "Docker", "FastAPI", "React"], jd["preferred_skills"]

jd2 = parse_job_description("""Data Engineer
Requirements:
- 2+ years of experience with Python and SQL
- Hands-on with Airflow, Kafka
- Experience with Docker is a plus
Preferred Skills: Kubernetes, Azure
Education: B.Tech in Computer Science, Machine Learning""")
assert set(jd2["required_skills"]) == {"Python", "SQL", "Airflow", "Kafka"}, jd2["required_skills"]
assert "Docker" in jd2["preferred_skills"] and "Kubernetes" in jd2["preferred_skills"]
assert "Machine Learning" not in jd2["all_skills"], "degree line must not create a skill requirement"

jd3 = parse_job_description("We need a developer who knows Python, React and Git. Docker is nice to have.")
assert set(jd3["required_skills"]) == {"Python", "React", "Git"} and jd3["preferred_skills"] == ["Docker"]

# ---------------------------------------------------------------- projects: one project = one entry
resume = parse_resume((ROOT / "sample_data" / "sample_resume.txt").read_text(encoding="utf-8"))
assert [p["name"] for p in resume["project_details"]] == ["DocuMind AI", "AI Smart Study Planner"]

multi = parse_resume("""RAVI KUMAR
ravi@x.com | +91 98765 43210
PROJECTS
DocuMind AI | Python, Streamlit | 2025
• Built a resume analyzer that parses PDFs and scores
  resumes using sentence embeddings
• Added semantic matching with MiniLM
AI Smart Study Planner
Developed a Streamlit app with SQLite for task planning
and progress tracking.
Face Attendance System (OpenCV, Python)
- Real-time face recognition using OpenCV
- Deployed on Flask
EDUCATION
B.Tech in AI & ML, XYZ University, 2022-2026
SKILLS
Languages: Python, Java, Kafka
Tools: Git, Docker
""")
names = [p["name"] for p in multi["project_details"]]
assert names == ["DocuMind AI", "AI Smart Study Planner", "Face Attendance System"], names
assert "embeddings" in multi["project_details"][0]["description"] and "MiniLM" in multi["project_details"][0]["description"]
assert "progress tracking" in multi["project_details"][1]["description"]
assert {"OpenCV", "Flask"} <= set(multi["project_details"][2]["skills"])
assert "Kafka" in multi["skills"] and "Machine Learning" not in multi["skills"]  # list item kept; degree title ignored
assert len(multi["projects"]) == 3

numbered = parse_resume("Projects\n1. Chatbot - Built a FAQ chatbot using NLP and Flask\n2. Sales Dashboard - Created Power BI dashboard on SQL data\nSkills\nPython")
assert [p["name"] for p in numbered["project_details"]] == ["Chatbot", "Sales Dashboard"]

# ---------------------------------------------------------------- projects that used to be split into several entries
from modules.resume_parser import parse_projects
def _names(txt): return [p["name"] for p in parse_projects(txt.split("\n"))]
assert _names("Smart Attendance System\nFace recognition based attendance\nUses OpenCV and Python\n"
              "Library Management\nJava Swing desktop app\nMySQL database") == ["Smart Attendance System", "Library Management"]
assert _names("E-Commerce Website    Jan 2024 - Mar 2024\nBuilt a full stack shop using React and Node\nIntegrated Stripe payments\n"
              "Weather App    Apr 2024\nMade a weather app using OpenWeather API") == ["E-Commerce Website", "Weather App"]
assert _names("Portfolio Website\nGitHub: github.com/x/portfolio\nResponsive portfolio built with HTML, CSS and JS\n"
              "Tic Tac Toe\nGitHub: github.com/x/ttt\nGame in Python") == ["Portfolio Website", "Tic Tac Toe"]
assert _names("Sentiment Analysis of Tweets using Machine\nLearning\nCollected tweets with Tweepy and trained a Naive\n"
              "Bayes classifier to classify sentiment with 87%\naccuracy.\nStudent Management System using Java and\nMySQL\n"
              "Created CRUD operations for students records.") == ["Sentiment Analysis of Tweets using Machine Learning",
                                                                     "Student Management System using Java and MySQL"]
assert _names("Hospital App - Flask\n- Patient records module\n- Billing module\nQuiz App - Django\n- Leaderboard") == ["Hospital App", "Quiz App"]
assert _names("Hotel Booking System\nTechnologies: Java, MySQL\nBuilt booking modules.\nStock Predictor\nTech Stack: Python, LSTM\n"
              "Predicts prices using LSTM.") == ["Hotel Booking System", "Stock Predictor"]

# ---------------------------------------------------------------- matching
m = compare_resume_to_jd(resume, jd)
assert 0 <= m["match_percentage"] <= 100
assert {"Python", "SQL", "Machine Learning", "Pandas", "Git"} <= set(m["matching_skills"])
assert set(m["missing_skills"]) == {"AWS", "Docker", "FastAPI", "React"}
assert m["partial_skills"] == ["Data Analysis"]
assert m["required_coverage"] > 85 and m["preferred_coverage"] == 0
assert [p["name"] for p in m["project_matches"]] and len(m["project_matches"]) == 2  # projects are scored individually
assert m["project_matches"][0]["relevance"] >= m["project_matches"][1]["relevance"]

other = parse_resume("Asha\nSKILLS\nJava, Spring Boot, MySQL\nPROJECTS\nBank App - Built a banking app in Java and Spring Boot with MySQL\n")
m_other = compare_resume_to_jd(other, jd)
assert m_other["match_percentage"] < m["match_percentage"] - 20, (m_other["match_percentage"], m["match_percentage"])

perfect = parse_resume("Z\nSKILLS\n" + ", ".join(jd["required_skills"] + jd["preferred_skills"]) + "\nPROJECTS\nTool - Built a Python machine learning data analysis tool with Pandas, SQL, Git, Docker, AWS, React and FastAPI")
assert compare_resume_to_jd(perfect, jd)["match_percentage"] > m["match_percentage"]

gap = build_skill_gap(resume["skills"], jd["required_skills"], m)
assert gap["high_priority"] == [] and "Data Analysis" in gap["medium_priority"] and "Docker" in gap["medium_priority"]

# ---------------------------------------------------------------- DOCX with a table
from docx import Document
d = Document()
d.add_paragraph("PROJECTS")
t = d.add_table(rows=1, cols=2)
t.cell(0, 0).text = "Chatbot - Built with Python and Flask"
t.cell(0, 1).text = "Dashboard - Power BI on SQL data"
buf = io.BytesIO(); d.save(buf)
class Up:  # minimal stand-in for Streamlit's UploadedFile
    name = "r.docx"
    def getvalue(self): return buf.getvalue()
text = extract_text(Up())
assert "Chatbot" in text and "Power BI" in text
print("All matching / parsing checks passed.")
